"""
The implied-vol file (decisions/0061, HANDOFF task 59): `bbg_impvol_values_only.xlsx`, two sheets of
series blocks (`IV_Weekly_v`, Bloomberg weekly on Fridays; `IV_Daily_v`, daily) and a probe
summary that is a record and not data.

Each series is a block of three columns, date, value and a blank. Rows 1 to 4 of a block are the
workbook's `series_id`, the security, the field and the overrides; data starts on row 5.

Pure functions on rows and frames. Values stay as delivered (implied vols in percentage points,
SKEW as an index level): no conversion, no interpolation, no fill (rule 5). Nothing here joins the
snapshot (rule 6).

The lab id. The workbook's ids look like JPMaQS tickers (`SPX_IV3M_10DP`: three capitals, an
underscore), and `loaders.jpmaqs.TICKER_RE` would classify them as JPMaQS. A series is therefore
named `<security>|<point>`, the part of the workbook id from `IV3M_` on, and no lab id can match
that pattern (rule 1). The workbook id is kept beside it.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd

from signal_lab.loaders.holdout import enforce_holdout, holdout_start
from signal_lab.loaders.invariants import InvariantViolation
from signal_lab.params import Params

ROLE = "implied_vol"
SHA256_PREFIX = "aeb90f8758e46fbd"
WEEKLY_SHEET = "IV_Weekly_v"
DAILY_SHEET = "IV_Daily_v"
HEADER_ROWS = 4
BLOCK_WIDTH = 3
POINT_MARKER = "IV3M_"
# decisions/0061 rule 4: a run of this many identical consecutive weekly values keeps its first
# week and marks the rest missing. Cowork's choice, not a measurement: longer than any run in the
# non-gold series (2 at most).
STALE_RUN_WEEKS = 6
FRIDAY = 4
LONG_COLUMNS = (
    "date",
    "series_id",
    "workbook_series_id",
    "security",
    "field",
    "overrides",
    "sheet",
    "value",
)


class RawFileChanged(RuntimeError):
    """The file on disk is not the one this loader was written against."""


@dataclass(frozen=True)
class ImpVolPanel:
    """`long` is the loaded panel; `stale` lists the weeks the stale rule removed."""

    long: pd.DataFrame
    stale: pd.DataFrame  # columns: series_id, date


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def check_file(path: Path) -> str:
    """The file's sha256, after checking it begins with the expected prefix."""
    digest = file_sha256(path)
    if not digest.startswith(SHA256_PREFIX):
        raise RawFileChanged(
            f"{Path(path).name} hashes to {digest[:16]}..., not {SHA256_PREFIX}...; the loader "
            f"was written against another file. Stop and tell the owner."
        )
    return digest


def lab_id(workbook_id: str, security: str) -> str:
    """Rule 1: `<security>|<point>`. `SKEW` is `SKEW Index|PX_LAST`."""
    if workbook_id == "SKEW":
        point = "PX_LAST"
    elif POINT_MARKER in workbook_id:
        point = workbook_id[workbook_id.index(POINT_MARKER) :]
    else:
        raise InvariantViolation(
            f"implied-vol series {workbook_id!r} has no {POINT_MARKER!r} and is not SKEW; "
            "no lab id can be made from it"
        )
    return f"{security}|{point}"


def _text(cell: Any) -> str:
    return "" if cell is None else str(cell).strip()


def read_blocks(rows: list[tuple], sheet: str) -> pd.DataFrame:
    """
    Sheet rows (header first) into a frame with a row per observation, `LONG_COLUMNS` less the
    holdout and week handling. A block with an empty header cell, or a value cell that is text (a
    Bloomberg error such as `#N/A N/A`), is refused with its series_id named. A blank value is
    a missing observation and is not kept.
    """
    if len(rows) <= HEADER_ROWS:
        raise InvariantViolation(f"{sheet}: fewer than {HEADER_ROWS + 1} rows")
    width = max(len(r) for r in rows)
    grid = [tuple(r) + (None,) * (width - len(r)) for r in rows]
    frames: list[pd.DataFrame] = []
    seen: set[str] = set()
    for c in range(0, width, BLOCK_WIDTH):
        header = [_text(grid[i][c]) for i in range(HEADER_ROWS)]
        body = [(r[c], r[c + 1] if c + 1 < width else None) for r in grid[HEADER_ROWS:]]
        if not any(header) and all(d is None and v is None for d, v in body):
            continue  # a trailing blank block
        wid = header[0] or f"<column {c + 1}>"
        if not all(header):
            raise InvariantViolation(
                f"{sheet}: the block at column {c + 1} ({wid}) has an empty header cell "
                f"(series_id, security, field, overrides = {header})"
            )
        dates, values = [], []
        for d, v in body:
            if d is None and v is None:
                continue
            if d is None:
                raise InvariantViolation(f"{sheet}: {wid} has a value with no date")
            if isinstance(v, str) or isinstance(v, bool):
                raise InvariantViolation(f"{sheet}: {wid} has a text value ({v!r}), not a number")
            if v is None:
                continue
            dates.append(pd.Timestamp(d).normalize())
            values.append(float(v))
        if len(set(dates)) != len(dates):
            raise InvariantViolation(f"{sheet}: {wid} repeats a date")
        security = header[1]
        if lab_id(wid, security) in seen:
            raise InvariantViolation(
                f"{sheet}: two series share the lab id {lab_id(wid, security)!r}"
            )
        seen.add(lab_id(wid, security))
        frames.append(
            pd.DataFrame(
                {
                    "date": dates,
                    "series_id": lab_id(wid, security),
                    "workbook_series_id": wid,
                    "security": security,
                    "field": header[2],
                    "overrides": header[3],
                    "sheet": sheet,
                    "value": values,
                }
            )
        )
    if not frames:
        raise InvariantViolation(f"{sheet}: no series blocks")
    out = pd.concat(frames, ignore_index=True)
    out["date"] = pd.to_datetime(out["date"])
    return out[list(LONG_COLUMNS)]


def to_fridays(frame: pd.DataFrame, sheet_kind: str, params: Params) -> pd.DataFrame:
    """
    Rules 2 and 3. The holdout first; then the weekly sheet is read as it is (every date must be a
    Friday), and a daily series becomes weekly by the last observation of each Monday-to-Friday
    week, dated by that week's Friday. A week whose Friday is on or after the holdout is dropped
    whole. Nothing is carried across weeks.
    """
    cut = holdout_start(params)
    frame = frame[frame["date"] < cut]
    if sheet_kind == "weekly":
        bad = frame[frame["date"].dt.weekday != FRIDAY]
        if len(bad):
            first = bad.iloc[0]
            raise InvariantViolation(
                f"{first['workbook_series_id']}: weekly date {first['date'].date()} is not a Friday"
            )
        return frame.reset_index(drop=True)
    if sheet_kind != "daily":
        raise ValueError(f"unknown sheet kind {sheet_kind!r}")
    weekend = frame[frame["date"].dt.weekday > FRIDAY]
    if len(weekend):
        first = weekend.iloc[0]
        raise InvariantViolation(
            f"{first['workbook_series_id']}: daily date {first['date'].date()} is a weekend"
        )
    # Holdout was cut on the day, but a week ending on or after it goes whole (rule 3).
    # The cut above can only leave a week whose Friday is later; those weeks are dropped below.
    frame = frame.copy()
    frame["_friday"] = frame["date"] + pd.to_timedelta(FRIDAY - frame["date"].dt.weekday, unit="D")
    frame = frame.sort_values(["series_id", "date"])
    last = frame.groupby(["series_id", "_friday"], sort=False).tail(1)
    last = last[last["_friday"] < cut].copy()
    last["date"] = last["_friday"]
    return last.drop(columns="_friday").reset_index(drop=True)


def drop_stale(frame: pd.DataFrame, run: int = STALE_RUN_WEEKS) -> ImpVolPanel:
    """
    Rule 4. In a series' weekly values a run of `run` or more consecutive Fridays holding the same
    value keeps its first week and marks the rest missing. A missing Friday ends a run.
    """
    frame = frame.sort_values(["series_id", "date"]).reset_index(drop=True)
    drop = pd.Series(False, index=frame.index)
    for _, g in frame.groupby("series_id", sort=False):
        same = (g["value"] == g["value"].shift()) & (g["date"].diff().dt.days == 7)
        run_id = (~same).cumsum()  # a run's first week is not `same`; each later week is
        in_long_run = g.groupby(run_id)["value"].transform("size") >= run
        drop.loc[g.index] = (same & in_long_run).to_numpy()
    removed = frame.loc[drop, ["series_id", "date"]].reset_index(drop=True)
    return ImpVolPanel(long=frame.loc[~drop].reset_index(drop=True), stale=removed)


def load_impvol(path: Path, params: Params) -> ImpVolPanel:
    """Read the workbook (hash-checked) into the weekly long panel, holdout and stale rule applied."""
    from openpyxl import load_workbook

    check_file(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet, kind in ((WEEKLY_SHEET, "weekly"), (DAILY_SHEET, "daily")):
            rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
            parts.append(to_fridays(read_blocks(rows, sheet), kind, params))
    finally:
        wb.close()
    long = pd.concat(parts, ignore_index=True)
    shared = set(parts[0]["series_id"]) & set(parts[1]["series_id"])
    if shared:
        raise InvariantViolation(f"two series share a lab id: {sorted(shared)[:5]}")
    panel = drop_stale(long)
    enforce_holdout(panel.long.set_index("date"), params, what="implied-vol panel")
    return panel
